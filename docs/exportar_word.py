"""Script para compilar DOCUMENTACION_MAESTRA_PITA.md en DOCUMENTACION_MAESTRA_PITA.docx
con diseño tipográfico formal para la Universidad Popular del Cesar.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_shading(cell, color_hex):
    """Aplica color de fondo a una celda."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Establece márgenes internos de una celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Aplica bordes finos a una tabla."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


def format_inlines(paragraph, text, font_name="Calibri", font_size=10.5, font_color=None):
    """Parsea negritas (**texto**), cursivas (*texto*) y código (`codigo`)."""
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = font_name
        run.font.size = Pt(font_size)
        if font_color:
            run.font.color.rgb = font_color

        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(font_size * 0.92)
            run.font.color.rgb = RGBColor(192, 41, 43)  # Rojo oscuro para código inline
        else:
            run.text = token


def main():
    md_file = os.path.join(os.path.dirname(__file__), "DOCUMENTACION_MAESTRA_PITA.md")
    docx_file = os.path.join(os.path.dirname(__file__), "DOCUMENTACION_MAESTRA_PITA.docx")

    with open(md_file, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Configuración de márgenes (2.5 cm = 0.984 pulgadas)
    for s in doc.sections:
        s.top_margin = Inches(0.98)
        s.bottom_margin = Inches(0.98)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.98)

    # Colores institucionales
    COLOR_NAVY = RGBColor(31, 78, 121)    # #1F4E79
    COLOR_TEAL = RGBColor(22, 58, 92)     # #163A5C
    COLOR_GREEN = RGBColor(46, 125, 50)   # #2E7D32
    COLOR_TEXT = RGBColor(40, 40, 40)
    COLOR_MUTED = RGBColor(100, 110, 120)

    # ─────────────────────────────────────────────────────────────
    # PORTADA FORMAL
    # ─────────────────────────────────────────────────────────────
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("UNIVERSIDAD POPULAR DEL CESAR\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(15)
    r_inst.bold = True
    r_inst.font.color.rgb = COLOR_NAVY

    r_fac = p_inst.add_run("FACULTAD DE INGENIERÍAS Y TECNOLÓGICAS\nDEPARTAMENTO DE INGENIERÍA DE SISTEMAS\n")
    r_fac.font.name = "Arial"
    r_fac.font.size = Pt(11)
    r_fac.font.color.rgb = COLOR_TEAL

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(70)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("DOCUMENTACIÓN MAESTRA DEL SISTEMA PITA\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    r_sub = p_title.add_run("PROGRAMA INTEGRADO DE TRANSACCIONES ACADÉMICAS\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_GREEN
    r_sub.bold = True

    r_tall = p_title.add_run("Taller 1 — Estructura de Datos (Listas en C++ y Python)")
    r_tall.font.name = "Calibri"
    r_tall.font.size = Pt(11.5)
    r_tall.font.color.rgb = COLOR_MUTED

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_after = Pt(120)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_docente = p_meta.add_run("Docente Titular:\nIng. Adith Pérez\n\n")
    r_docente.font.name = "Calibri"
    r_docente.font.size = Pt(11)
    r_docente.bold = True

    r_lugar = p_meta.add_run("Valledupar, Cesar — Colombia\nSeptiembre de 2026")
    r_lugar.font.name = "Calibri"
    r_lugar.font.size = Pt(10.5)
    r_lugar.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PARSEO Y VOLCADO DEL CONTENIDO
    # ─────────────────────────────────────────────────────────────
    i = 0
    in_code_block = False
    code_lines = []

    # Ignoramos la cabecera repetida del markdown hasta el primer encabezado real
    started = False

    while i < len(lines):
        line = lines[i].rstrip()

        # Inicio del contenido tras la metadata inicial
        if line.startswith("## Control del Documento"):
            started = True

        if not started:
            i += 1
            continue

        # Bloque de código
        if line.startswith("```"):
            if in_code_block:
                # Termina bloque de código
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.3)
                r = p.add_run("\n".join(code_lines))
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Encabezados
        if line.startswith("# "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(15)
            r.bold = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        elif line.startswith("## "):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(12.5)
            r.bold = True
            r.font.color.rgb = COLOR_TEAL
            i += 1
            continue

        elif line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(11.5)
            r.bold = True
            r.font.color.rgb = COLOR_GREEN
            i += 1
            continue

        elif line.startswith("#### "):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = COLOR_TEAL
            i += 1
            continue

        # Tablas en Markdown
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                # Parsear filas
                rows_data = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    rows_data.append(cells)

                # Si la segunda fila es separador de alineación (contiene dashes)
                if len(rows_data) > 1 and all(set(c).issubset({'-', ':', ' '}) for c in rows_data[1]):
                    header = rows_data[0]
                    data_rows = rows_data[2:]
                else:
                    header = rows_data[0]
                    data_rows = rows_data[1:]

                col_count = len(header)
                tbl = doc.add_table(rows=len(data_rows) + 1, cols=col_count)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="D1D5DB")

                # Fila de cabecera
                hdr_cells = tbl.rows[0].cells
                for c_idx, title in enumerate(header):
                    if c_idx < len(hdr_cells):
                        cell = hdr_cells[c_idx]
                        set_cell_shading(cell, "1F4E79")
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        p_c = cell.paragraphs[0]
                        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        format_inlines(p_c, title, font_name="Arial", font_size=9.5, font_color=RGBColor(255, 255, 255))
                        p_c.runs[0].bold = True if p_c.runs else False

                # Filas de datos
                for r_idx, drow in enumerate(data_rows):
                    row_cells = tbl.rows[r_idx + 1].cells
                    bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, val in enumerate(drow):
                        if c_idx < len(row_cells):
                            cell = row_cells[c_idx]
                            set_cell_shading(cell, bg_color)
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            p_c = cell.paragraphs[0]
                            format_inlines(p_c, val, font_name="Calibri", font_size=9.5, font_color=COLOR_TEXT)

                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
            continue

        # Línea horizontal divisoria
        if line.startswith("---"):
            i += 1
            continue

        # Viñetas
        if line.startswith("- ") or line.startswith("* "):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inlines(p, text, font_name="Calibri", font_size=10.5, font_color=COLOR_TEXT)
            i += 1
            continue

        # Numeradas
        m_num = re.match(r'^(\d+)\.\s+(.*)', line)
        if m_num:
            text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inlines(p, text, font_name="Calibri", font_size=10.5, font_color=COLOR_TEXT)
            i += 1
            continue

        # Párrafo ordinario
        if line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            format_inlines(p, line, font_name="Calibri", font_size=10.5, font_color=COLOR_TEXT)

        i += 1

    doc.save(docx_file)
    print(f"Documento Word generado exitosamente: {docx_file}")


if __name__ == "__main__":
    main()
